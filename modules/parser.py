# ============================================
# parser.py - 문서에서 텍스트를 추출하는 모듈
# ============================================
# 목적: 다양한 형식의 문서를 읽어서 텍스트와 메타데이터를 추출한다
#
# 지원 형식:
#   - PDF: pdfplumber로 페이지별 텍스트 추출
#   - DOCX: python-docx로 문단별 텍스트 추출
#   - PPTX: python-pptx로 슬라이드별 텍스트 추출
#   - TXT: 파이썬 내장 함수로 텍스트 읽기
#   - XLSX: openpyxl로 시트별 텍스트 추출
#
# [변경] v2.0 - 멀티 포맷 지원 추가
#   기존 PDF 전용 → PDF/DOCX/PPTX/TXT 지원
# ============================================

import pdfplumber  # PDF 파일을 열고 텍스트를 추출하는 라이브러리
import os  # 파일 경로, 폴더 탐색 등 운영체제 관련 기능
import re  # 표 데이터 필터링용

def parse_single_pdf(file_path):
    """
    단일 PDF 파일에서 페이지별 텍스트와 메타데이터를 추출하는 함수

    파라미터:
        file_path (str): PDF 파일의 전체 경로 (예: "data/pdfs/06_Guidebook.pdf")

    반환값:
        list[dict]: 페이지별 정보를 담은 딕셔너리 리스트
            - "text": 해당 페이지에서 추출한 텍스트
            - "metadata": 파일명, 페이지 번호 등 부가 정보
    """
    documents = []  # 추출된 페이지 정보를 저장할 리스트

    file_name = os.path.basename(file_path)  # 전체 경로에서 파일명만 추출 (예: "06_Guidebook.pdf")

    try:  # PDF 파일 열기 시도 (깨진 파일이나 잘못된 경로 대비)
        with pdfplumber.open(file_path) as pdf:  # PDF 파일을 열어서 pdf 객체 생성
            total_pages = len(pdf.pages)  # 전체 페이지 수 확인
            print(f"  📄 {file_name}: 총 {total_pages}페이지")  # 처리 상황 출력

            for page_num, page in enumerate(pdf.pages):  # 각 페이지를 순회 (0부터 시작)
                text = page.extract_text()

                if text and text.strip():  # 텍스트가 존재하고 빈 문자열이 아닌 경우만 처리
                    clean_text = clean_extracted_text(text)  # 추출된 텍스트 정제

                    # 200자 미만 페이지 필터링 (표지, 빈 페이지 등 노이즈 제거)
                    # 표지/빈 페이지는 "가이드북", 페이지 번호 정도만 있어서 검색 노이즈가 됨
                    if len(clean_text.strip()) < 200:
                        # print(f"    ⚠️  {page_num + 1}페이지: {len(clean_text.strip())}자 (노이즈 페이지 → 건너뜀)")
                        continue  # 이 페이지는 건너뛰기

                    document = {  # 페이지 정보를 딕셔너리로 구성
                        "text": clean_text,  # 정제된 텍스트
                        "metadata": {  # 메타데이터 (나중에 출처 표시에 사용)
                            "source": file_name,  # 원본 파일명
                            "page": page_num + 1,  # 페이지 번호 (1부터 시작)
                            "total_pages": total_pages,  # 전체 페이지 수
                        }
                    }
                    documents.append(document)  # 리스트에 추가
                else:  # 텍스트가 없는 페이지 (이미지만 있는 페이지 등)
                    print(f"    ⚠️  {page_num + 1}페이지: 텍스트 없음 (이미지 페이지일 수 있음)")

    except Exception as e:  # PDF 파일 열기 또는 파싱 실패 시
        print(f"  ❌ {file_name} 파싱 실패: {e}")  # 에러 메시지 출력

    return documents  # 추출된 전체 페이지 정보 반환


def clean_extracted_text(text):
    """
    PDF에서 추출한 텍스트를 정제하는 함수

    파라미터:
        text (str): PDF에서 추출한 원본 텍스트

    반환값:
        str: 정제된 텍스트

    처리 내용:
        - 줄 앞뒤 공백 제거
        - 빈 줄이 3개 이상 연속되면 2개로 축소
    """
    lines = text.split("\n")  # 텍스트를 줄 단위로 분리
    cleaned_lines = []  # 정제된 줄을 저장할 리스트
    empty_count = 0  # 연속 빈 줄 개수를 추적하는 카운터

    for line in lines:  # 각 줄을 순회
        stripped = line.strip()  # 줄 앞뒤의 공백 제거

        # 표/행렬 데이터 라인 필터링
        # d1(1x1), d414(256x320) 같은 데이터 구조 표기를 제거
        # 이런 데이터가 청크에 포함되면 임베딩 품질이 저하되어 검색 실패 유발
        if stripped and re.search(r'd\d+\(\d+x\d+\)', stripped):
            continue  # 이 줄은 건너뛰기

        if stripped:  # 내용이 있는 줄인 경우
            cleaned_lines.append(stripped)  # 정제된 줄 추가
            empty_count = 0  # 빈 줄 카운터 초기화
        else:  # 빈 줄인 경우
            empty_count += 1  # 빈 줄 카운터 증가
            if empty_count <= 2:  # 빈 줄이 2개 이하면 유지 (문단 구분 보존)
                cleaned_lines.append("")  # 빈 줄 추가

    result = "\n".join(cleaned_lines)  # 줄들을 다시 하나의 문자열로 합침
    return result  # 정제된 텍스트 반환


def parse_all_pdfs(pdf_directory):
    """
    지정된 폴더 안의 모든 PDF 파일을 파싱하는 함수

    파라미터:
        pdf_directory (str): PDF 파일들이 있는 폴더 경로 (예: "data/pdfs")

    반환값:
        list[dict]: 모든 PDF의 모든 페이지 정보를 담은 리스트
    """
    all_documents = []  # 전체 문서 정보를 저장할 리스트

    if not os.path.exists(pdf_directory):  # 폴더가 없으면
        print(f"❌ 폴더를 찾을 수 없습니다: {pdf_directory}")  # 에러 메시지 출력
        return all_documents  # 빈 리스트 반환

    # 폴더 안에서 .pdf 파일만 필터링하여 정렬
    pdf_files = sorted([  # 파일명 기준으로 정렬
        f for f in os.listdir(pdf_directory)  # 폴더 안의 모든 파일 목록
        if f.lower().endswith(".pdf")  # 확장자가 .pdf인 파일만 (대소문자 무시)
    ])

    if not pdf_files:  # PDF 파일이 하나도 없으면
        print(f"⚠️  {pdf_directory} 폴더에 PDF 파일이 없습니다.")  # 경고 메시지
        return all_documents  # 빈 리스트 반환

    print(f"\n📂 {len(pdf_files)}개의 PDF 파일을 발견했습니다.\n")  # 발견된 파일 수 출력

    for pdf_file in pdf_files:  # 각 PDF 파일을 순회
        file_path = os.path.join(pdf_directory, pdf_file)  # 폴더 경로 + 파일명 = 전체 경로
        documents = parse_single_pdf(file_path)  # 해당 PDF 파싱 실행
        all_documents.extend(documents)  # 결과를 전체 리스트에 추가

    print(f"\n✅ 총 {len(all_documents)}개의 페이지를 추출했습니다.\n")  # 최종 결과 출력
    return all_documents  # 전체 문서 정보 반환


# ============================================
# [신규] DOCX 파싱 함수
# ============================================
# 사용 라이브러리: python-docx
# 설치: pip install python-docx
# ============================================

def parse_single_docx(file_path):
    """
    단일 DOCX 파일에서 텍스트와 메타데이터를 추출하는 함수

    DOCX는 PDF와 달리 "페이지" 개념이 없으므로,
    문서 전체를 하나의 텍스트로 추출한다.
    (렌더링 엔진 없이는 페이지 구분이 불가능)

    파라미터:
        file_path (str): DOCX 파일의 전체 경로

    반환값:
        list[dict]: 문서 정보 리스트 (보통 1개 항목)
            - "text": 추출된 텍스트
            - "metadata": 파일명 등 부가 정보
    """
    from docx import Document as DocxDocument  # python-docx 라이브러리

    documents = []
    file_name = os.path.basename(file_path)

    try:
        doc = DocxDocument(file_path)  # DOCX 파일 열기

        # 모든 문단(paragraph)의 텍스트를 줄바꿈으로 합침
        # doc.paragraphs: 문서의 모든 문단 객체 리스트
        # p.text: 해당 문단의 텍스트 (서식 제거된 순수 텍스트)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        text = "\n".join(paragraphs)

        if text and len(text.strip()) >= 50:  # 50자 미만이면 빈 문서로 간주
            clean_text = clean_extracted_text(text)

            document = {
                "text": clean_text,
                "metadata": {
                    "source": file_name,
                    "page": 1,  # DOCX는 페이지 구분 불가 → 1로 고정
                    "total_pages": 1,
                    "file_type": "docx",  # [신규] 파일 형식 구분용
                }
            }
            documents.append(document)
            print(f"  📄 {file_name}: {len(clean_text)}자 추출")
        else:
            print(f"  ⚠️  {file_name}: 텍스트 없음 또는 너무 짧음")

    except Exception as e:
        print(f"  ❌ {file_name} 파싱 실패: {e}")

    return documents


# ============================================
# [신규] PPTX 파싱 함수
# ============================================
# 사용 라이브러리: python-pptx
# 설치: pip install python-pptx
# ============================================

def parse_single_pptx(file_path):
    """
    단일 PPTX 파일에서 슬라이드별 텍스트와 메타데이터를 추출하는 함수

    PPTX는 슬라이드 단위로 구분되므로,
    각 슬라이드를 별도의 문서(페이지)로 처리한다.
    PDF의 페이지 개념과 유사.

    파라미터:
        file_path (str): PPTX 파일의 전체 경로

    반환값:
        list[dict]: 슬라이드별 정보 리스트
            - "text": 해당 슬라이드에서 추출한 텍스트
            - "metadata": 파일명, 슬라이드 번호 등
    """
    from pptx import Presentation  # python-pptx 라이브러리

    documents = []
    file_name = os.path.basename(file_path)

    try:
        prs = Presentation(file_path)  # PPTX 파일 열기
        total_slides = len(prs.slides)
        print(f"  📄 {file_name}: 총 {total_slides}슬라이드")

        for slide_num, slide in enumerate(prs.slides):  # 각 슬라이드 순회
            texts = []

            # 슬라이드 안의 모든 도형(shape)에서 텍스트 추출
            # shape: 텍스트 상자, 제목, 본문, 표 등 슬라이드 위의 모든 요소
            for shape in slide.shapes:
                if shape.has_text_frame:  # 텍스트가 있는 도형만 처리
                    # text_frame: 도형 안의 텍스트 영역
                    # paragraphs: 텍스트 영역 안의 문단들
                    for paragraph in shape.text_frame.paragraphs:
                        para_text = paragraph.text.strip()
                        if para_text:
                            texts.append(para_text)

            # 슬라이드의 모든 텍스트를 합침
            slide_text = "\n".join(texts)

            # 텍스트가 너무 짧은 슬라이드는 건너뜀 (제목만 있는 슬라이드 등)
            if slide_text and len(slide_text.strip()) >= 20:
                document = {
                    "text": slide_text,
                    "metadata": {
                        "source": file_name,
                        "page": slide_num + 1,  # 슬라이드 번호 (1부터 시작)
                        "total_pages": total_slides,
                        "file_type": "pptx",
                    }
                }
                documents.append(document)

    except Exception as e:
        print(f"  ❌ {file_name} 파싱 실패: {e}")

    return documents


# ============================================
# [신규] TXT 파싱 함수
# ============================================

def parse_single_txt(file_path):
    """
    단일 TXT 파일에서 텍스트와 메타데이터를 추출하는 함수

    TXT는 가장 단순한 형식으로, 파일 내용을 그대로 읽으면 된다.
    인코딩은 UTF-8을 기본으로 하고, 실패 시 CP949(한국어)를 시도한다.

    파라미터:
        file_path (str): TXT 파일의 전체 경로

    반환값:
        list[dict]: 문서 정보 리스트 (보통 1개 항목)
            - "text": 파일 전체 텍스트
            - "metadata": 파일명 등 부가 정보
    """
    documents = []
    file_name = os.path.basename(file_path)

    try:
        # UTF-8로 먼저 시도, 실패 시 CP949(한국어 Windows 인코딩) 시도
        text = None
        for encoding in ["utf-8", "cp949", "euc-kr"]:
            try:
                with open(file_path, "r", encoding=encoding) as f:
                    text = f.read()
                break  # 성공하면 루프 탈출
            except UnicodeDecodeError:
                continue  # 다음 인코딩으로 시도

        if text is None:
            print(f"  ❌ {file_name}: 인코딩을 판별할 수 없습니다.")
            return documents

        if text and len(text.strip()) >= 50:
            clean_text = clean_extracted_text(text)

            document = {
                "text": clean_text,
                "metadata": {
                    "source": file_name,
                    "page": 1,  # TXT는 페이지 개념 없음 → 1로 고정
                    "total_pages": 1,
                    "file_type": "txt",
                }
            }
            documents.append(document)
            print(f"  📄 {file_name}: {len(clean_text)}자 추출")
        else:
            print(f"  ⚠️  {file_name}: 텍스트 없음 또는 너무 짧음")

    except Exception as e:
        print(f"  ❌ {file_name} 파싱 실패: {e}")

    return documents


# ============================================
# [신규] XLSX 파싱 함수
# ============================================
# 사용 라이브러리: openpyxl
# 설치: pip install openpyxl
#
# 주의: Excel은 숫자/표 데이터 위주라 임베딩 검색이 잘 안 될 수 있음.
#       텍스트가 포함된 Excel(용어 정의표, FAQ 목록 등)에 적합.
# ============================================

def parse_single_xlsx(file_path):
    """
    단일 XLSX 파일에서 시트별 텍스트와 메타데이터를 추출하는 함수

    각 시트를 별도의 문서(페이지)로 처리한다.
    각 행의 셀 값을 탭으로 구분하여 한 줄로 합치고,
    모든 행을 줄바꿈으로 연결한다.

    예시:
        | 용어 | 정의 |
        | RAG  | 검색 증강 생성 |
        → "용어\t정의\nRAG\t검색 증강 생성"

    파라미터:
        file_path (str): XLSX 파일의 전체 경로

    반환값:
        list[dict]: 시트별 정보 리스트
            - "text": 해당 시트에서 추출한 텍스트
            - "metadata": 파일명, 시트 이름 등
    """
    from openpyxl import load_workbook  # Excel 파일 읽기 라이브러리

    documents = []
    file_name = os.path.basename(file_path)

    try:
        # data_only=True: 수식 대신 계산된 값을 읽음
        # 예: =SUM(A1:A10) → 수식이 아니라 결과값 55를 가져옴
        wb = load_workbook(file_path, data_only=True)
        total_sheets = len(wb.sheetnames)
        print(f"  📄 {file_name}: 총 {total_sheets}개 시트")

        for sheet_idx, sheet_name in enumerate(wb.sheetnames):
            ws = wb[sheet_name]  # 시트 객체 가져오기
            rows_text = []

            for row in ws.iter_rows(values_only=True):
                # 각 셀의 값을 문자열로 변환, None은 빈 문자열로
                # 탭(\t)으로 구분하여 한 줄로 합침
                cells = [str(cell) if cell is not None else "" for cell in row]
                row_text = "\t".join(cells).strip()

                # 완전히 빈 행은 건너뜀
                if row_text and row_text.replace("\t", ""):
                    rows_text.append(row_text)

            # 모든 행을 줄바꿈으로 합침
            sheet_text = "\n".join(rows_text)

            if sheet_text and len(sheet_text.strip()) >= 20:
                document = {
                    "text": sheet_text,
                    "metadata": {
                        "source": file_name,
                        "page": sheet_idx + 1,  # 시트 번호 (1부터 시작)
                        "total_pages": total_sheets,
                        "sheet_name": sheet_name,  # 시트 이름도 저장
                        "file_type": "xlsx",
                    }
                }
                documents.append(document)
                print(f"    📊 시트 '{sheet_name}': {len(sheet_text)}자 추출")
            else:
                print(f"    ⚠️  시트 '{sheet_name}': 텍스트 없음 또는 너무 짧음")

    except Exception as e:
        print(f"  ❌ {file_name} 파싱 실패: {e}")

    return documents


# ============================================
# [신규] 파일 형식 자동 감지 + 파싱 통합 함수
# ============================================
# 지원 확장자 목록 (app.py에서도 사용)
SUPPORTED_EXTENSIONS = [".pdf", ".docx", ".pptx", ".txt", ".xlsx"]


def parse_single_file(file_path):
    """
    파일 확장자를 감지하여 적절한 파서를 호출하는 통합 함수

    파라미터:
        file_path (str): 파일의 전체 경로

    반환값:
        list[dict]: 파싱된 문서 정보 리스트
            - PDF: 페이지별 리스트
            - PPTX: 슬라이드별 리스트
            - DOCX/TXT: 보통 1개 항목
    """
    # 파일 확장자를 소문자로 추출 (예: ".PDF" → ".pdf")
    _, ext = os.path.splitext(file_path)
    ext = ext.lower()

    if ext == ".pdf":
        return parse_single_pdf(file_path)
    elif ext == ".docx":
        return parse_single_docx(file_path)
    elif ext == ".pptx":
        return parse_single_pptx(file_path)
    elif ext == ".txt":
        return parse_single_txt(file_path)
    elif ext == ".xlsx":
        return parse_single_xlsx(file_path)
    else:
        print(f"  ❌ 지원하지 않는 파일 형식: {ext}")
        return []


# ============================================
# 직접 실행 시 테스트 코드
# 실행 방법: python modules/parser.py
# ============================================
if __name__ == "__main__":
    PDF_DIR = "data/pdfs"  # PDF 파일이 저장된 폴더 경로

    results = parse_all_pdfs(PDF_DIR)  # 모든 PDF 파싱 실행

    # 결과 미리보기: 처음 3개 페이지의 텍스트 앞부분만 출력
    for i, doc in enumerate(results[:3]):  # 최대 3개만 출력
        print(f"--- 문서 {i+1} ---")  # 구분선
        print(f"출처: {doc['metadata']['source']} / {doc['metadata']['page']}페이지")  # 출처 정보
        text = doc['text']
        if len(text) > 500:  # 500자보다 길면
            print(f"내용 미리보기: {text[:500]}...")  # 잘랐으니 ... 표시
        else:  # 500자 이하면
            print(f"내용 미리보기: {text}")  # 전체 출력, ... 없음
        print()  # 빈 줄